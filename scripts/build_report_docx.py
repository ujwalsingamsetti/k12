"""Script to compile the complete, publication-grade B.E. Final Year Capstone Project Report
for Context-Aware Automated Grading of Written Exams as an editable Microsoft Word document (.docx).

Strict adherence to institutional guidelines:
- Paper Size: A4 (8.27" x 11.69")
- Margins: Left 1.25", Right 1.0", Top 1.0", Bottom 1.0"
- Main Chapter Headings: Arial 14pt Bold, Centered
- Subheadings: Arial 12pt Bold, Left-aligned
- Body Text: Arial 12pt, 1.5 line spacing, Justified, flush-left (first_line_indent = 0)
- Abstract: Single paragraph containing 15-20 lines (19 lines), Arial 12pt, 1.5 line spacing
- Clean booktabs tables: pure white background everywhere, subtle horizontal rules (1.0pt top/bottom, 0.75pt mid), zero vertical borders
- Embedded institutional banners (assets/image1.jpeg and assets/image2.jpeg)
- Embedded Figures 1 through 9 from docs/figures (Figures 10, 11, and UI screenshots excluded)
- References: Alphabetical order by author name (minimum 12, 20 included), no website URLs
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from report_data import (
    STUDENT_CONFIG, STUDENT_1_NAME, STUDENT_1_REG, STUDENT_2_NAME, STUDENT_2_REG,
    GUIDE_NAME, GUIDE_DESIGNATION, HOD_NAME, HOD_DESIGNATION,
    DEPARTMENT_NAME, SCHOOL_NAME, UNIVERSITY_NAME, UNIVERSITY_SUB,
    UNIVERSITY_ACCREDITATION, LOCATION, MONTH_YEAR, PROJECT_TITLE,
    CERT_TEXT, DECL_TEXT, ACK_PARAS, ABSTRACT_TEXT, KEYWORDS_TEXT,
    ABBREVIATIONS, TOC_PAGE_1, TOC_PAGE_2, LIST_OF_FIGURES, LIST_OF_TABLES
)
import report_content_ch1_ch3 as ch1_3
import report_content_ch4_ch6 as ch4_6

DOCX_OUTPUT_PATH = "K12_Automated_Grading_Final_Report.docx"

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ASSETS_DIR = os.path.join(BASE_DIR, "assets")
FIGURES_DIR = os.path.join(BASE_DIR, "k12-answer-evaluator", "docs", "figures")


def set_cell_background(cell, fill_hex="FFFFFF"):
    """Sets background color of a docx table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    """Sets padding/margins for a docx table cell in twips."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)


def set_table_booktabs_borders(table):
    """Applies clean booktabs horizontal rules:
    - 1.0pt top rule (#334155)
    - 1.0pt bottom rule (#334155)
    - 0.75pt mid rule below header row (#334155)
    - Zero vertical rules or inner vertical lines.
    """
    tblPr = table._tbl.tblPr
    existing_borders = tblPr.find(qn('w:tblBorders'))
    if existing_borders is not None:
        tblPr.remove(existing_borders)
    
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="8" w:space="0" w:color="334155"/>'
        f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="334155"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="none"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

    for cell in table.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="334155"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(tcBorders)


def set_table_borderless(table):
    """Applies completely borderless style for TOC."""
    tblPr = table._tbl.tblPr
    existing_borders = tblPr.find(qn('w:tblBorders'))
    if existing_borders is not None:
        tblPr.remove(existing_borders)
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="none"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def add_footer_page_number(run):
    """Adds a dynamic PAGE field code to a footer paragraph run."""
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)


def configure_section_page_numbering(section, fmt="decimal", start_num=None):
    """Sets section page numbering format (lowerRoman or decimal)."""
    sectPr = section._sectPr
    for child in list(sectPr):
        if child.tag.endswith('pgNumType'):
            sectPr.remove(child)
    pgNumType = OxmlElement('w:pgNumType')
    pgNumType.set(qn('w:fmt'), fmt)
    if start_num is not None:
        pgNumType.set(qn('w:start'), str(start_num))
    sectPr.append(pgNumType)


def add_centered_heading(doc, text):
    """Adds a bold, centered chapter heading in Arial 14pt."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    return p


def add_section_heading(doc, text):
    """Adds a bold, left-aligned section heading in Arial 12pt."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    return p


def add_subsection_heading(doc, text):
    """Adds a bold italic, left-aligned subsection heading in Arial 12pt."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = RGBColor(30, 41, 59)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    return p


def clean_xml_tags(text):
    """Strips ReportLab-style pseudo-HTML tags from text for python-docx."""
    import re
    cleaned = re.sub(r'<[^>]+>', '', text)
    cleaned = cleaned.replace('&lt;', '<').replace('&gt;', '>').replace('&amp;', '&').replace('&rarr;', '->')
    return cleaned


def add_body_paragraph(doc, text):
    """Adds a justified body paragraph in Arial 12pt with 1.5 line spacing, flush left."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(clean_xml_tags(text))
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(30, 41, 59)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Inches(0.0)
    return p


def add_bullet_paragraph(doc, text):
    """Adds an indented bullet paragraph."""
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(clean_xml_tags(text))
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(30, 41, 59)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    return p


def add_figure(doc, img_path, caption_text, width=Inches(5.5)):
    """Embeds an image followed by a centered italic caption."""
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.add_run().add_picture(img_path, width=width)

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_cap = p_cap.add_run(caption_text)
        run_cap.font.name = 'Arial'
        run_cap.font.size = Pt(9.5)
        run_cap.font.italic = True
        run_cap.font.color.rgb = RGBColor(51, 65, 85)
        p_cap.paragraph_format.space_before = Pt(3)
        p_cap.paragraph_format.space_after = Pt(10)


def add_academic_table(doc, caption_text, headers, data, col_widths=None):
    """Adds a clean booktabs-style academic table with caption above."""
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_cap = p_cap.add_run(caption_text)
    run_cap.font.name = 'Arial'
    run_cap.font.size = Pt(10)
    run_cap.font.bold = True
    run_cap.font.color.rgb = RGBColor(15, 23, 42)
    p_cap.paragraph_format.space_before = Pt(8)
    p_cap.paragraph_format.space_after = Pt(3)
    p_cap.paragraph_format.keep_with_next = True

    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_booktabs_borders(table)

    # Format Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = clean_xml_tags(title)
        set_cell_background(hdr_cells[i], "FFFFFF")
        set_cell_margins(hdr_cells[i], top=80, bottom=80, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = 'Arial'
            run.font.size = Pt(9)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)

    # Format Data Rows
    for r_idx, row in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = clean_xml_tags(str(val))
            set_cell_background(row_cells[c_idx], "FFFFFF")
            set_cell_margins(row_cells[c_idx], top=60, bottom=60, left=100, right=100)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = 'Arial'
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(30, 41, 59)

    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = width

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table


def build_docx_report():
    doc = Document()

    # ═════════════════════════════════════════════════════════════════════
    # SECTION 1: COVER / TITLE PAGE (Suppressed Footer)
    # ═════════════════════════════════════════════════════════════════════
    sec1 = doc.sections[0]
    sec1.page_width = Inches(8.27)
    sec1.page_height = Inches(11.69)
    sec1.left_margin = Inches(1.25)
    sec1.right_margin = Inches(1.0)
    sec1.top_margin = Inches(1.0)
    sec1.bottom_margin = Inches(1.0)
    sec1.different_first_page_header_footer = True

    img1_path = os.path.join(ASSETS_DIR, "image1.jpeg")
    img2_path = os.path.join(ASSETS_DIR, "image2.jpeg")

    if os.path.exists(img1_path):
        p_img1 = doc.add_paragraph()
        p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img1.add_run().add_picture(img1_path, width=Inches(4.8))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(PROJECT_TITLE)
    run.font.name = 'Arial'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(8)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("PROJECT REPORT")
    run.font.name = 'Arial'
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = RGBColor(51, 65, 85)
    p.paragraph_format.space_after = Pt(6)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Submitted in partial fulfillment of the requirements for the award of\nBachelor of Engineering degree in Computer Science and Engineering")
    run.font.name = 'Arial'
    run.font.size = Pt(10)
    run.font.italic = True
    p.paragraph_format.space_after = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"By\n{STUDENT_1_NAME} (Reg. No – {STUDENT_1_REG})\n{STUDENT_2_NAME} (Reg. No – {STUDENT_2_REG})")
    run.font.name = 'Arial'
    run.font.size = Pt(10.5)
    run.font.bold = True
    p.paragraph_format.space_after = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"{DEPARTMENT_NAME}\n{SCHOOL_NAME}")
    run.font.name = 'Arial'
    run.font.size = Pt(11)
    run.font.bold = True
    p.paragraph_format.space_after = Pt(8)

    if os.path.exists(img2_path):
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.add_run().add_picture(img2_path, width=Inches(4.8))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"{UNIVERSITY_NAME}\n{UNIVERSITY_SUB}\n{UNIVERSITY_ACCREDITATION}\n{LOCATION}\n\n{MONTH_YEAR}")
    run.font.name = 'Arial'
    run.font.size = Pt(10)
    p.paragraph_format.space_before = Pt(6)

    # ═════════════════════════════════════════════════════════════════════
    # SECTION 2: FRONT MATTER (Lowercase Roman Numerals ii to x)
    # ═════════════════════════════════════════════════════════════════════
    sec2 = doc.add_section()
    sec2.left_margin = Inches(1.25)
    sec2.right_margin = Inches(1.0)
    sec2.top_margin = Inches(1.0)
    sec2.bottom_margin = Inches(1.0)
    sec2.header.is_linked_to_previous = False
    sec2.footer.is_linked_to_previous = False
    configure_section_page_numbering(sec2, fmt="lowerRoman", start_num=2)

    f_para = sec2.footer.paragraphs[0]
    f_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f_run = f_para.add_run()
    f_run.font.name = 'Arial'
    f_run.font.size = Pt(10)
    add_footer_page_number(f_run)

    # --- BONAFIDE CERTIFICATE (Page ii) ---
    p_dep = doc.add_paragraph()
    p_dep.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_dep = p_dep.add_run(DEPARTMENT_NAME)
    run_dep.font.name = 'Arial'
    run_dep.font.size = Pt(12)
    run_dep.font.bold = True
    p_dep.paragraph_format.space_before = Pt(10)
    p_dep.paragraph_format.space_after = Pt(4)

    add_centered_heading(doc, "BONAFIDE CERTIFICATE")
    add_body_paragraph(doc, CERT_TEXT)

    doc.add_paragraph().paragraph_format.space_after = Pt(28)

    t_sig = doc.add_table(rows=1, cols=2)
    t_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    c1 = t_sig.rows[0].cells[0].paragraphs[0]
    c1.add_run(f"Internal Guide\n\n\n\n{GUIDE_NAME}\n{GUIDE_DESIGNATION}").font.bold = True
    c2 = t_sig.rows[0].cells[1].paragraphs[0]
    c2.add_run(f"Head of the Department\n\n\n\n{HOD_NAME}\n{HOD_DESIGNATION}").font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(24)

    add_body_paragraph(doc, "Submitted for Project Viva Voce Examination held on ____________________")
    t_viva = doc.add_table(rows=1, cols=2)
    t_viva.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_viva.rows[0].cells[0].paragraphs[0].add_run("Internal Examiner").font.bold = True
    t_viva.rows[0].cells[1].paragraphs[0].add_run("External Examiner").font.bold = True

    doc.add_page_break()

    # --- DECLARATION (Page iii) ---
    add_centered_heading(doc, "DECLARATION")
    add_body_paragraph(doc, DECL_TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(36)
    add_body_paragraph(doc, f"DATE: _____________\nPLACE: Chennai\n\nSIGNATURE OF THE CANDIDATES:\n1. ____________________________ ({STUDENT_1_NAME})\n2. ____________________________ ({STUDENT_2_NAME})")

    doc.add_page_break()

    # --- ACKNOWLEDGEMENT (Page iv) ---
    add_centered_heading(doc, "ACKNOWLEDGEMENT")
    for p_txt in ACK_PARAS:
        add_body_paragraph(doc, p_txt)

    doc.add_page_break()

    # --- ABSTRACT (Page v) ---
    add_centered_heading(doc, "ABSTRACT")
    add_body_paragraph(doc, ABSTRACT_TEXT)
    add_body_paragraph(doc, f"Keywords: {KEYWORDS_TEXT}")

    doc.add_page_break()

    # --- TABLE OF CONTENTS (Pages vi - vii) ---
    add_centered_heading(doc, "TABLE OF CONTENTS")
    t_toc1 = doc.add_table(rows=len(TOC_PAGE_1) + 1, cols=3)
    t_toc1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borderless(t_toc1)
    t_toc1.rows[0].cells[0].paragraphs[0].add_run("CHAPTER NO.").font.bold = True
    t_toc1.rows[0].cells[1].paragraphs[0].add_run("TITLE").font.bold = True
    t_toc1.rows[0].cells[2].paragraphs[0].add_run("PAGE NO.").font.bold = True
    t_toc1.rows[0].cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    for r_idx, (ch_no, title, pg, is_bold) in enumerate(TOC_PAGE_1):
        r_cells = t_toc1.rows[r_idx + 1].cells
        r_cells[0].text = ch_no
        r_cells[1].text = title
        r_cells[2].text = pg
        r_cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if is_bold:
            for c in r_cells:
                for run in c.paragraphs[0].runs:
                    run.font.bold = True

    doc.add_page_break()

    add_centered_heading(doc, "TABLE OF CONTENTS (CONTINUED)")
    t_toc2 = doc.add_table(rows=len(TOC_PAGE_2) + 1, cols=3)
    t_toc2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borderless(t_toc2)
    t_toc2.rows[0].cells[0].paragraphs[0].add_run("CHAPTER NO.").font.bold = True
    t_toc2.rows[0].cells[1].paragraphs[0].add_run("TITLE").font.bold = True
    t_toc2.rows[0].cells[2].paragraphs[0].add_run("PAGE NO.").font.bold = True
    t_toc2.rows[0].cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    for r_idx, (ch_no, title, pg, is_bold) in enumerate(TOC_PAGE_2):
        r_cells = t_toc2.rows[r_idx + 1].cells
        r_cells[0].text = ch_no
        r_cells[1].text = title
        r_cells[2].text = pg
        r_cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if is_bold:
            for c in r_cells:
                for run in c.paragraphs[0].runs:
                    run.font.bold = True

    doc.add_page_break()

    # --- LIST OF ABBREVIATIONS (Page viii) ---
    add_centered_heading(doc, "LIST OF ABBREVIATIONS")
    t_abb = doc.add_table(rows=len(ABBREVIATIONS) + 1, cols=2)
    t_abb.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_booktabs_borders(t_abb)
    t_abb.rows[0].cells[0].paragraphs[0].add_run("ABBREVIATION").font.bold = True
    t_abb.rows[0].cells[1].paragraphs[0].add_run("EXPANDED DEFINITION").font.bold = True
    for r_idx, (ab, exp) in enumerate(ABBREVIATIONS):
        t_abb.rows[r_idx + 1].cells[0].text = ab
        t_abb.rows[r_idx + 1].cells[1].text = exp

    doc.add_page_break()

    # --- LIST OF FIGURES (Page ix - Separate Page!) ---
    add_centered_heading(doc, "LIST OF FIGURES")
    t_lf = doc.add_table(rows=len(LIST_OF_FIGURES) + 1, cols=3)
    t_lf.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_booktabs_borders(t_lf)
    t_lf.rows[0].cells[0].paragraphs[0].add_run("FIGURE NO.").font.bold = True
    t_lf.rows[0].cells[1].paragraphs[0].add_run("FIGURE NAME").font.bold = True
    t_lf.rows[0].cells[2].paragraphs[0].add_run("PAGE NO.").font.bold = True
    t_lf.rows[0].cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r_idx, (fno, fname, fpg) in enumerate(LIST_OF_FIGURES):
        t_lf.rows[r_idx + 1].cells[0].text = fno
        t_lf.rows[r_idx + 1].cells[1].text = fname
        t_lf.rows[r_idx + 1].cells[2].text = fpg
        t_lf.rows[r_idx + 1].cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_page_break()

    # --- LIST OF TABLES (Page x - Separate Page!) ---
    add_centered_heading(doc, "LIST OF TABLES")
    t_lt = doc.add_table(rows=len(LIST_OF_TABLES) + 1, cols=3)
    t_lt.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_booktabs_borders(t_lt)
    t_lt.rows[0].cells[0].paragraphs[0].add_run("TABLE NO.").font.bold = True
    t_lt.rows[0].cells[1].paragraphs[0].add_run("TABLE NAME").font.bold = True
    t_lt.rows[0].cells[2].paragraphs[0].add_run("PAGE NO.").font.bold = True
    t_lt.rows[0].cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r_idx, (tno, tname, tpg) in enumerate(LIST_OF_TABLES):
        t_lt.rows[r_idx + 1].cells[0].text = tno
        t_lt.rows[r_idx + 1].cells[1].text = tname
        t_lt.rows[r_idx + 1].cells[2].text = tpg
        t_lt.rows[r_idx + 1].cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_page_break()

    # ═════════════════════════════════════════════════════════════════════
    # SECTION 3: CHAPTER 1 ONWARDS (Arabic Numerals 1 to 50)
    # ═════════════════════════════════════════════════════════════════════
    sec3 = doc.add_section()
    sec3.left_margin = Inches(1.25)
    sec3.right_margin = Inches(1.0)
    sec3.top_margin = Inches(1.0)
    sec3.bottom_margin = Inches(1.0)
    sec3.header.is_linked_to_previous = False
    sec3.footer.is_linked_to_previous = False
    configure_section_page_numbering(sec3, fmt="decimal", start_num=1)

    f_para3 = sec3.footer.paragraphs[0]
    f_para3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f_run3 = f_para3.add_run()
    f_run3.font.name = 'Arial'
    f_run3.font.size = Pt(10)
    add_footer_page_number(f_run3)

    # ─────────────────────────────────────────────────────────────────────
    # CHAPTER 1: INTRODUCTION
    # ─────────────────────────────────────────────────────────────────────
    add_centered_heading(doc, "CHAPTER 1\nINTRODUCTION")
    for sec in ch1_3.CH1_SECTIONS:
        if sec['level'] == 2:
            add_section_heading(doc, sec['heading'])
        else:
            add_subsection_heading(doc, sec['heading'])
        for p in sec['paragraphs']:
            add_body_paragraph(doc, p)

    # Table 1.1
    add_academic_table(doc, "Table 1.1: Systematic Comparison of Information Retrieval Paradigms in Exam Evaluation",
                       ch1_3.TABLE_1_1_HEADERS, ch1_3.TABLE_1_1_DATA)

    # Fig 1.1
    fig1_path = os.path.join(FIGURES_DIR, "fig1_overall_system_architecture.png")
    add_figure(doc, fig1_path, "Fig 1.1: Master System Architecture and End-to-End Evaluation Workflow")

    for sec in ch1_3.CH1_SECTIONS_CONT:
        add_section_heading(doc, sec['heading'])
        for p in sec['paragraphs']:
            add_body_paragraph(doc, p)

    doc.add_page_break()

    # ─────────────────────────────────────────────────────────────────────
    # CHAPTER 2: LITERATURE SURVEY
    # ─────────────────────────────────────────────────────────────────────
    add_centered_heading(doc, "CHAPTER 2\nLITERATURE SURVEY")
    for sec in ch1_3.CH2_SECTIONS_P8:
        if sec['level'] == 2:
            add_section_heading(doc, sec['heading'])
        else:
            add_subsection_heading(doc, sec['heading'])
        for p in sec['paragraphs']:
            add_body_paragraph(doc, p)

    for sec in ch1_3.CH2_SECTIONS_P9:
        add_subsection_heading(doc, sec['heading'])
        for p in sec['paragraphs']:
            add_body_paragraph(doc, p)

    add_section_heading(doc, ch1_3.CH2_SECTIONS_P10[0]['heading'])
    for p in ch1_3.CH2_SECTIONS_P10[0]['paragraphs']:
        add_body_paragraph(doc, p)

    # Table 2.1 Part 1 & Part 2
    add_academic_table(doc, "Table 2.1: Literature Review Matrix (14 State-of-the-Art Peer-Reviewed Papers - Part 1)",
                       ch1_3.TABLE_2_1A_HEADERS, ch1_3.TABLE_2_1A_DATA)
    add_academic_table(doc, "Table 2.1: Literature Review Matrix (14 State-of-the-Art Peer-Reviewed Papers - Part 2)",
                       ch1_3.TABLE_2_1B_HEADERS, ch1_3.TABLE_2_1B_DATA)

    # Fig 2.1
    fig7_path = os.path.join(FIGURES_DIR, "fig7_empirical_performance_benchmarks.png")
    add_figure(doc, fig7_path, "Fig 2.1: Empirical Performance & Baseline Benchmark Comparison")

    # Table 2.2
    add_academic_table(doc, "Table 2.2: Deep Dive Architectural Feature Benchmark: EduEval Baseline vs. Proposed System",
                       ch1_3.TABLE_2_2_HEADERS, ch1_3.TABLE_2_2_DATA)

    for sec in ch1_3.CH2_SECTIONS_CONT2:
        add_section_heading(doc, sec['heading'])
        for p in sec['paragraphs']:
            add_body_paragraph(doc, p)

    doc.add_page_break()

    # ─────────────────────────────────────────────────────────────────────
    # CHAPTER 3: AIM AND SCOPE
    # ─────────────────────────────────────────────────────────────────────
    add_centered_heading(doc, "CHAPTER 3\nAIM AND SCOPE OF THE PRESENT INVESTIGATION")
    for sec in ch1_3.CH3_SECTIONS:
        add_section_heading(doc, sec['heading'])
        for p in sec['paragraphs']:
            add_body_paragraph(doc, p)

    # Table 3.3
    add_academic_table(doc, "Table 3.3: Target Operational User Personas and Workflow Specifications",
                       ch1_3.TABLE_3_3_HEADERS, ch1_3.TABLE_3_3_DATA)

    for sec in ch1_3.CH3_SECTIONS_CONT:
        add_section_heading(doc, sec['heading'])
        for p in sec['paragraphs']:
            add_body_paragraph(doc, p)

    # Table 3.1
    add_academic_table(doc, "Table 3.1: Feasibility Study & Implementation Risk Assessment Matrix",
                       ch1_3.TABLE_3_1_HEADERS, ch1_3.TABLE_3_1_DATA)

    for sec in ch1_3.CH3_SECTIONS_CONT2:
        add_section_heading(doc, sec['heading'])
        for p in sec['paragraphs']:
            add_body_paragraph(doc, p)

    # Table 3.2
    add_academic_table(doc, "Table 3.2: Hardware, Operating System, and Software Stack Specifications",
                       ch1_3.TABLE_3_2_HEADERS, ch1_3.TABLE_3_2_DATA)

    # Fig 3.1
    fig5_path = os.path.join(FIGURES_DIR, "fig5_multipage_answer_collation.png")
    add_figure(doc, fig5_path, "Fig 3.1: Multi-Page Answer Collation and Ingestion Sequence")

    doc.add_page_break()

    # ─────────────────────────────────────────────────────────────────────
    # CHAPTER 4: MATERIALS AND METHODS
    # ─────────────────────────────────────────────────────────────────────
    add_centered_heading(doc, "CHAPTER 4\nMATERIALS AND METHODS; ALGORITHMS USED")
    
    # 4.1 Master Blueprint
    add_section_heading(doc, ch4_6.CH4_SECTIONS[0]['heading'])
    for p in ch4_6.CH4_SECTIONS[0]['paragraphs']:
        add_body_paragraph(doc, p)

    # 4.2 Module 1 & Fig 4.1
    add_section_heading(doc, ch4_6.CH4_SECTIONS[1]['heading'])
    for p in ch4_6.CH4_SECTIONS[1]['paragraphs']:
        add_body_paragraph(doc, p)
    fig2_path = os.path.join(FIGURES_DIR, "fig2_vision_preprocessing_pipeline.png")
    add_figure(doc, fig2_path, "Fig 4.1: Optical Character Recognition & Vision Preprocessing Pipeline")

    # 4.3 Module 2 & Algorithm 1
    add_section_heading(doc, ch4_6.CH4_SECTIONS[2]['heading'])
    for p in ch4_6.CH4_SECTIONS[2]['paragraphs']:
        add_body_paragraph(doc, p)
    add_body_paragraph(doc, ch4_6.ALGORITHM_1_TEXT)

    # 4.3.1 & 4.3.2
    for sec in ch4_6.CH4_SECTIONS_MOD2_SUB:
        add_subsection_heading(doc, sec['heading'])
        for p in sec['paragraphs']:
            add_body_paragraph(doc, p)

    # 4.4 Module 3 & Fig 4.2
    add_section_heading(doc, ch4_6.CH4_SECTIONS_MOD3[0]['heading'])
    for p in ch4_6.CH4_SECTIONS_MOD3[0]['paragraphs']:
        add_body_paragraph(doc, p)
    fig8_path = os.path.join(FIGURES_DIR, "fig8_database_er_and_vector_topology.png")
    add_figure(doc, fig8_path, "Fig 4.2: Relational Schema & Qdrant HNSW Vector Topology")

    # 4.4.1 & Table 4.1
    add_subsection_heading(doc, ch4_6.CH4_SECTIONS_MOD3[1]['heading'])
    for p in ch4_6.CH4_SECTIONS_MOD3[1]['paragraphs']:
        add_body_paragraph(doc, p)
    add_academic_table(doc, "Table 4.1: PostgreSQL Relational Schema & Qdrant Payload Collections",
                       ch4_6.TABLE_4_1_HEADERS, ch4_6.TABLE_4_1_DATA)

    # 4.5 Module 4 & Fig 4.3
    add_section_heading(doc, ch4_6.CH4_SECTIONS_MOD4[0]['heading'])
    for p in ch4_6.CH4_SECTIONS_MOD4[0]['paragraphs']:
        add_body_paragraph(doc, p)
    fig4_path = os.path.join(FIGURES_DIR, "fig4_dual_route_grading_rubric.png")
    add_figure(doc, fig4_path, "Fig 4.3: Dual-Route Automated Grading Rubric & Decision Flow")

    # 4.5.1 & Table 4.2 & Algorithm 2
    add_subsection_heading(doc, ch4_6.CH4_SECTIONS_MOD4[1]['heading'])
    for p in ch4_6.CH4_SECTIONS_MOD4[1]['paragraphs']:
        add_body_paragraph(doc, p)
    add_academic_table(doc, "Table 4.2: Few-Shot Prompt Templates & Multi-Attribute Scoring Parameters",
                       ch4_6.TABLE_4_2_HEADERS, ch4_6.TABLE_4_2_DATA)
    add_body_paragraph(doc, ch4_6.ALGORITHM_2_TEXT)

    # 4.5.2 & Fig 4.4
    add_subsection_heading(doc, ch4_6.CH4_SECTIONS_MOD4_GOV[0]['heading'])
    for p in ch4_6.CH4_SECTIONS_MOD4_GOV[0]['paragraphs']:
        add_body_paragraph(doc, p)
    fig6_path = os.path.join(FIGURES_DIR, "fig6_human_in_the_loop_governance.png")
    add_figure(doc, fig6_path, "Fig 4.4: Human-in-the-Loop Governance & Teacher Verification Workflow")

    # 4.6 Module 5 & Fig 4.5
    add_section_heading(doc, ch4_6.CH4_SECTIONS_MOD5[0]['heading'])
    for p in ch4_6.CH4_SECTIONS_MOD5[0]['paragraphs']:
        add_body_paragraph(doc, p)
    fig3_path = os.path.join(FIGURES_DIR, "fig3_vector_retrieval_rag_pipeline.png")
    add_figure(doc, fig3_path, "Fig 4.5: Vector Retrieval & Curricular RAG Pipeline")

    # 4.6.1 & Table 4.3a & Algorithm 3
    add_subsection_heading(doc, ch4_6.CH4_SECTIONS_MOD5[1]['heading'])
    for p in ch4_6.CH4_SECTIONS_MOD5[1]['paragraphs']:
        add_body_paragraph(doc, p)
    add_academic_table(doc, "Table 4.3a: Numerical Step-by-Step Walkthrough of Keyword Density Re-ranking",
                       ch4_6.TABLE_4_3A_HEADERS, ch4_6.TABLE_4_3A_DATA)
    add_body_paragraph(doc, ch4_6.ALGORITHM_3_TEXT)

    # 4.7 Module 6 & 4.7.1, 4.7.2
    for sec in ch4_6.CH4_SECTIONS_MOD6:
        if sec['level'] == 2:
            add_section_heading(doc, sec['heading'])
        else:
            add_subsection_heading(doc, sec['heading'])
        for p in sec['paragraphs']:
            add_body_paragraph(doc, p)

    # 4.8 Module 7 & Fig 4.6
    add_section_heading(doc, ch4_6.CH4_SECTIONS_MOD7[0]['heading'])
    for p in ch4_6.CH4_SECTIONS_MOD7[0]['paragraphs']:
        add_body_paragraph(doc, p)
    add_subsection_heading(doc, ch4_6.CH4_SECTIONS_MOD7[1]['heading'])
    for p in ch4_6.CH4_SECTIONS_MOD7[1]['paragraphs']:
        add_body_paragraph(doc, p)
    fig9_path = os.path.join(FIGURES_DIR, "fig9_mathematical_algorithmic_formulations.png")
    add_figure(doc, fig9_path, "Fig 4.6: Mathematical Formulations & Multi-Criteria Scoring Formulations")

    # 4.9 Cost Analysis & Table 4.3 & 4.10 Web Architecture
    add_section_heading(doc, ch4_6.CH4_SECTIONS_MOD8_9[0]['heading'])
    for p in ch4_6.CH4_SECTIONS_MOD8_9[0]['paragraphs']:
        add_body_paragraph(doc, p)
    add_academic_table(doc, "Table 4.3: Enterprise Total Cost of Ownership (TCO) & Infrastructure Profiling",
                       ch4_6.TABLE_4_3_HEADERS, ch4_6.TABLE_4_3_DATA)
    add_section_heading(doc, ch4_6.CH4_SECTIONS_MOD10[0]['heading'])
    for p in ch4_6.CH4_SECTIONS_MOD10[0]['paragraphs']:
        add_body_paragraph(doc, p)

    doc.add_page_break()

    # ─────────────────────────────────────────────────────────────────────
    # CHAPTER 5: RESULTS AND DISCUSSION
    # ─────────────────────────────────────────────────────────────────────
    add_centered_heading(doc, "CHAPTER 5\nRESULTS AND DISCUSSION, PERFORMANCE ANALYSIS")
    for sec in ch4_6.CH5_SECTIONS:
        add_section_heading(doc, sec['heading'])
        for p in sec['paragraphs']:
            add_body_paragraph(doc, p)

    # Table 5.1
    add_academic_table(doc, "Table 5.1: Multi-Criteria Question Grading Accuracy Benchmark Across Question Complexities",
                       ch4_6.TABLE_5_1_HEADERS, ch4_6.TABLE_5_1_DATA)

    for sec in ch4_6.CH5_SECTIONS_CONT:
        add_section_heading(doc, sec['heading'])
        for p in sec['paragraphs']:
            add_body_paragraph(doc, p)

    # Table 5.2
    add_academic_table(doc, "Table 5.2: Retrieval Precision, Faithfulness, and Hallucination Suppression Metrics",
                       ch4_6.TABLE_5_2_HEADERS, ch4_6.TABLE_5_2_DATA)

    # 5.5 & Table 5.3
    add_section_heading(doc, ch4_6.CH5_SECTIONS_CONT2[0]['heading'])
    for p in ch4_6.CH5_SECTIONS_CONT2[0]['paragraphs']:
        add_body_paragraph(doc, p)
    add_academic_table(doc, "Table 5.3: Ablation Study Quantifying the Empirical Impact of Core Architectural Modules",
                       ch4_6.TABLE_5_3_HEADERS, ch4_6.TABLE_5_3_DATA)

    # 5.6 & 5.7 & Table 5.4
    add_section_heading(doc, ch4_6.CH5_SECTIONS_CONT3[0]['heading'])
    for p in ch4_6.CH5_SECTIONS_CONT3[0]['paragraphs']:
        add_body_paragraph(doc, p)
    add_section_heading(doc, ch4_6.CH5_SECTIONS_CONT3[1]['heading'])
    for p in ch4_6.CH5_SECTIONS_CONT3[1]['paragraphs']:
        add_body_paragraph(doc, p)
    add_academic_table(doc, "Table 5.4: End-to-End Latency Telemetry Waterfall Breakdown Across 100 Test Submissions",
                       ch4_6.TABLE_5_4_HEADERS, ch4_6.TABLE_5_4_DATA)

    # 5.8 & Table 5.5 & 5.9 & Table 5.6
    add_section_heading(doc, ch4_6.CH5_SECTIONS_CONT4[0]['heading'])
    for p in ch4_6.CH5_SECTIONS_CONT4[0]['paragraphs']:
        add_body_paragraph(doc, p)
    add_academic_table(doc, "Table 5.5: Concurrent Throughput and P95 Tail Latency Scaling Benchmark",
                       ch4_6.TABLE_5_5_HEADERS, ch4_6.TABLE_5_5_DATA)
    add_section_heading(doc, ch4_6.CH5_SECTIONS_CONT5[0]['heading'])
    for p in ch4_6.CH5_SECTIONS_CONT5[0]['paragraphs']:
        add_body_paragraph(doc, p)
    add_academic_table(doc, "Table 5.6: Robustness Benchmark Under OCR Noise, Lexical Typos, and Diagram Dropout",
                       ch4_6.TABLE_5_6_HEADERS, ch4_6.TABLE_5_6_DATA)

    # 5.10 & Table 5.7
    add_section_heading(doc, ch4_6.CH5_SECTIONS_CONT6[0]['heading'])
    for p in ch4_6.CH5_SECTIONS_CONT6[0]['paragraphs']:
        add_body_paragraph(doc, p)
    add_academic_table(doc, "Table 5.7: Cross-Framework Comparative Benchmark Against Industry Systems",
                       ch4_6.TABLE_5_7_HEADERS, ch4_6.TABLE_5_7_DATA)

    # 5.11 & 5.12
    for sec in ch4_6.CH5_SECTIONS_CONT7:
        add_section_heading(doc, sec['heading'])
        for p in sec['paragraphs']:
            add_body_paragraph(doc, p)

    doc.add_page_break()

    # ─────────────────────────────────────────────────────────────────────
    # CHAPTER 6: SUMMARY AND CONCLUSIONS
    # ─────────────────────────────────────────────────────────────────────
    add_centered_heading(doc, "CHAPTER 6\nSUMMARY, CONCLUSIONS AND FUTURE WORK")
    for sec in ch4_6.CH6_SECTIONS:
        add_section_heading(doc, sec['heading'])
        for p in sec['paragraphs']:
            add_body_paragraph(doc, p)

    doc.add_page_break()

    # ─────────────────────────────────────────────────────────────────────
    # REFERENCES & BIBLIOGRAPHY
    # ─────────────────────────────────────────────────────────────────────
    add_centered_heading(doc, "REFERENCES")
    for ref_str in ch4_6.REFERENCES_LIST:
        add_body_paragraph(doc, ref_str)

    # Table R.1
    add_academic_table(doc, "Table R.1: Systematic Architectural Mapping of Literature References to System Subsystems",
                       ch4_6.TABLE_R_1_HEADERS, ch4_6.TABLE_R_1_DATA)

    # Table R.2
    add_academic_table(doc, "Table R.2: Cross-Reference Correlation Between Research Objectives and Benchmark Citations",
                       ch4_6.TABLE_R_2_HEADERS, ch4_6.TABLE_R_2_DATA)

    print(f"Compiling publication-grade report DOCX to {DOCX_OUTPUT_PATH}...")
    doc.save(DOCX_OUTPUT_PATH)
    print("DOCX Build complete.")


if __name__ == "__main__":
    build_docx_report()
